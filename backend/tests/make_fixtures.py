"""In-memory test fixture builders for attachment formats (PRD §54).

Builds bytes directly rather than committing binary files to the repo, so
fixtures stay text-diffable and cannot silently go stale. DOCX uses
python-docx (already a runtime dependency) as a writer. PDF is hand-rolled
byte-for-byte (a minimal-but-valid single/multi-page PDF with a real text
content stream) rather than adding a PDF-writing dependency such as
reportlab, which is not in the PRD's approved dependency list (PRD §39) and
would only ever be used for test fixtures.
"""

from __future__ import annotations

import io
import zipfile

import docx
import pypdf


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def build_docx_bytes(
    paragraphs: list[str] | None = None, table_rows: list[list[str]] | None = None
) -> bytes:
    document = docx.Document()
    if paragraphs:
        for text in paragraphs:
            document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=0, cols=len(table_rows[0]))
        for row_values in table_rows:
            row = table.add_row()
            for cell, value in zip(row.cells, row_values, strict=False):
                cell.text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def build_corrupted_docx_bytes() -> bytes:
    return build_docx_bytes(paragraphs=["SELECT 1 FROM DUAL"])[:80]


def build_docx_with_macro_bytes() -> bytes:
    base = build_docx_bytes(paragraphs=["SELECT * FROM T A WHERE A.X=1"])
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("word/vbaProject.bin", b"fake-macro-bytes")
    return out.getvalue()


def build_highly_compressed_docx_bytes(uncompressed_size: int = 5_000_000) -> bytes:
    """A docx-shaped zip with one entry that compresses at an extreme ratio
    (a long run of one byte) — exercises the zip-bomb compression-ratio
    guard without needing a multi-hundred-MB fixture on disk."""
    base = build_docx_bytes(paragraphs=["placeholder"])
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("word/huge.bin", b"A" * uncompressed_size, compress_type=zipfile.ZIP_DEFLATED)
    return out.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def _pdf_escape(text: str) -> str:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def build_pdf_bytes(pages: list[str]) -> bytes:
    objects: list[bytes] = []

    def add_object(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    font_num = add_object(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    content_nums = []
    for page_text in pages:
        lines = page_text.split("\n") or [""]
        y = 750
        parts = ["BT", "/F1 12 Tf"]
        for line in lines:
            parts.append(f"1 0 0 1 50 {y} Tm ({_pdf_escape(line)}) Tj")
            y -= 16
        parts.append("ET")
        stream = "\n".join(parts).encode("latin-1", errors="replace")
        body = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
        content_nums.append(add_object(body))

    pages_num = add_object(b"")  # placeholder, filled in once page numbers are known
    page_nums = []
    for content_num in content_nums:
        body = (
            b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>"
        ) % (pages_num, font_num, content_num)
        page_nums.append(add_object(body))

    kids = " ".join(f"{n} 0 R" for n in page_nums).encode("ascii")
    objects[pages_num - 1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, len(page_nums))

    catalog_num = add_object(b"<< /Type /Catalog /Pages %d 0 R >>" % pages_num)

    buf = io.BytesIO()
    buf.write(b"%PDF-1.4\n")
    offsets = [0] * (len(objects) + 1)
    for i, body in enumerate(objects, start=1):
        offsets[i] = buf.tell()
        buf.write(b"%d 0 obj\n" % i)
        buf.write(body)
        buf.write(b"\nendobj\n")

    xref_offset = buf.tell()
    buf.write(b"xref\n")
    buf.write(b"0 %d\n" % (len(objects) + 1))
    buf.write(b"0000000000 65535 f \n")
    for i in range(1, len(objects) + 1):
        buf.write(b"%010d 00000 n \n" % offsets[i])
    buf.write(b"trailer\n<< /Size %d /Root %d 0 R >>\n" % (len(objects) + 1, catalog_num))
    buf.write(b"startxref\n%d\n" % xref_offset)
    buf.write(b"%%EOF")
    return buf.getvalue()


def build_scanned_pdf_bytes() -> bytes:
    """A structurally valid PDF with no real text content — stands in for a
    scanned image PDF for the purposes of the "no text layer" heuristic."""
    return build_pdf_bytes([""])


def build_encrypted_pdf_bytes(text: str = "SELECT * FROM T A WHERE A.X=1", password: str = "secret") -> bytes:
    reader = pypdf.PdfReader(io.BytesIO(build_pdf_bytes([text])))
    writer = pypdf.PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt(password)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Text encodings (PRD §54 TXT cases)
# ---------------------------------------------------------------------------
def build_utf8_bytes(text: str) -> bytes:
    return text.encode("utf-8")


def build_utf8_bom_bytes(text: str) -> bytes:
    return b"\xef\xbb\xbf" + text.encode("utf-8")


def build_big5_bytes(text: str) -> bytes:
    return text.encode("big5")
